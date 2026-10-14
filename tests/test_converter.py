from datetime import datetime
from io import BytesIO
import pytest
from openpyxl import Workbook, load_workbook
from p6toolkit.converter import (parse_relationship_details, convert_rows, convert_workbook,
                                 to_import_workbook, ConverterError)

HEADER = ["Activity ID", "Activity Status", "Activity Name", "WBS", "WBS Name", "Orig Dur",
          "Actual Duration", "Rem Dur", "Start", "Finish", "Actual Start", "Actual Finish",
          "Primary Constraint", "Primary Constraint Date", "Predecessor Details", "Successor Details"]


def row(aid, status, name, od, rem, start, finish, a_s=None, a_f=None, pred=None, succ=None, wbs="P.1"):
    return dict(zip(HEADER, [aid, status, name, wbs, "W", od, 0, rem, start, finish, a_s, a_f,
                             None, None, pred, succ]))


def sample():
    return [
        row("M1", "Completed", "Start", 0, 0, "01-Sep-25 A", None, a_s=datetime(2025, 9, 1), succ="A1: FS"),
        row("A1", "Completed", "Design", 5, 0, "02-Sep-25 A", "08-Sep-25 A", datetime(2025, 9, 2),
            datetime(2025, 9, 8), pred="M1: FS", succ="A2: SS 2, A2: FF"),
        row("A2", "Not Started", "Build", 10, 10, datetime(2026, 1, 1), datetime(2026, 1, 14),
            pred="A1: SS 2, A1: FF", succ="M2: FS"),
        row("M2", "Not Started", "Finish", 0, 0, None, datetime(2026, 1, 14), pred="A2: FS"),
        dict(zip(HEADER, ["A2", "Not Started", None, None] + [None] * 10 + ["A1", None])),  # overflow
    ]


def test_parse_tokens():
    toks = parse_relationship_details("A1980: FS, M1210: SS 57 ,B: FF -3, bad, C: XX")
    assert [(t.activity, t.kind, t.lag, t.ok) for t in toks] == [
        ("A1980", "FS", 0.0, True), ("M1210", "SS", 57.0, True), ("B", "FF", -3.0, True),
        ("bad", None, None, False), ("C", "XX", 0.0, False)]
    assert parse_relationship_details(None) == []


def test_convert_counts_types_and_milestones():
    r = convert_rows(sample())
    assert r.activity_count == 4 and r.dropped_rows == 1
    assert r.relationship_count == 4
    assert r.relationship_types == {"FS": 2, "SS": 1, "FF": 1}
    assert r.lagged_relationships == 1
    assert r.activity_types == {"Start Milestone": 1, "Task Dependent": 2, "Finish Milestone": 1}
    assert [i.check for i in r.issues] == ["Dropped row (export cell overflow)"]
    pred = {(x[3], x[4], x[5]): (x[6], x[7], x[8]) for x in r.taskpred_rows}
    assert pred[("A1", "A2", "SS")] == ("Completed", "Not Started", 2)


def test_mismatch_and_missing_ids_reported():
    rows = sample()[:4]
    rows[1]["Successor Details"] = "A2: SS 2"          # FF missing on successor side
    rows[3]["Predecessor Details"] = "A2: FS, ZZ9: FS"  # unknown predecessor
    checks = [i.check for i in convert_rows(rows).issues]
    assert "Relationship only in Predecessor Details" in checks
    assert "Predecessor not found in data" in checks


def test_missing_columns_rejected():
    with pytest.raises(ConverterError):
        convert_rows([{"Activity ID": "A"}])


def test_workbook_round_trip():
    wb = Workbook(); ws = wb.active; ws.title = "data"; ws.append(HEADER)
    for r in sample():
        ws.append([r[h] for h in HEADER])
    buf = BytesIO(); wb.save(buf)
    result = convert_workbook(buf.getvalue())
    out = load_workbook(BytesIO(to_import_workbook(result)))
    assert out.sheetnames == ["TASK", "TASKPRED", "Validation"]
    assert out["TASK"]["A1"].value == "task_code" and out["TASK"]["A2"].value == "Activity ID"
    assert out["TASKPRED"].max_row == 2 + 4
    assert out["TASK"]["E3"].value == "Start Milestone"


def test_missing_sheet_rejected():
    wb = Workbook(); wb.active.title = "other"; buf = BytesIO(); wb.save(buf)
    with pytest.raises(ConverterError):
        convert_workbook(buf.getvalue())
