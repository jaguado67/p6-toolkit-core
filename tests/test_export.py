from datetime import datetime
from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from p6toolkit.export import dataframe_to_xlsx, safe_sheet_name


def test_round_trip_values_types_and_header():
    df = pd.DataFrame({"Activity ID": ["A1", "A2"], "Total Float (h)": [8.0, None],
                       "Date": pd.to_datetime([datetime(2026, 10, 1, 8), None]),
                       "Note": ["ok\x07bell", None]})
    ws = load_workbook(BytesIO(dataframe_to_xlsx(df, "Float"))).active
    assert ws.title == "Float"
    assert [c.value for c in ws[1]] == ["Activity ID", "Total Float (h)", "Date", "Note"]
    assert ws["A2"].value == "A1" and ws["B2"].value == 8 and ws["B3"].value is None
    assert ws["C2"].value == datetime(2026, 10, 1, 8) and ws["C2"].number_format == "yyyy-mm-dd hh:mm"
    assert ws["D2"].value == "okbell"            # control character removed, not a corrupt file
    assert ws.freeze_panes == "A2"


def test_sheet_names_are_valid():
    assert safe_sheet_name("xer_field_changes.csv/[x]") == "xer_field_changes.csv__x_"
    assert len(safe_sheet_name("a" * 50)) == 31


def test_empty_frame_exports():
    ws = load_workbook(BytesIO(dataframe_to_xlsx(pd.DataFrame(columns=["A"]), "E"))).active
    assert ws.max_row == 1
