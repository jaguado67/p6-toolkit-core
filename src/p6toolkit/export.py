"""Excel export for any result table (Streamlit-independent, testable)."""
from __future__ import annotations
from io import BytesIO
import re
import pandas as pd

EXCEL_MAX_ROWS = 1_048_576
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")   # rejected by the XLSX format
_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")


def safe_sheet_name(name: str) -> str:
    clean = _SHEET_BAD.sub("_", name).strip("'") or "Sheet1"
    return clean[:31]


def dataframe_to_xlsx(df: pd.DataFrame, sheet_name: str = "Data") -> bytes:
    """One formatted sheet: bold frozen header, filter, widths, dates as real Excel dates."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    if len(df) + 1 > EXCEL_MAX_ROWS:
        raise ValueError(f"{len(df):,} rows exceed the Excel limit of {EXCEL_MAX_ROWS - 1:,}")
    out = df.copy()
    for col in out.columns:
        if not (pd.api.types.is_numeric_dtype(out[col]) or pd.api.types.is_datetime64_any_dtype(out[col])):
            out[col] = out[col].map(lambda v: _ILLEGAL.sub("", v) if isinstance(v, str) else v)
    buf = BytesIO()
    name = safe_sheet_name(sheet_name)
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        out.to_excel(writer, index=False, sheet_name=name)
        ws = writer.sheets[name]
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        if len(out):
            ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(out.columns, start=1):
            sample = [str(col)] + [str(v) for v in out[col].head(200).tolist()]
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, max(map(len, sample)) + 2))
            if pd.api.types.is_datetime64_any_dtype(out[col]):
                for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                    for c in cell:
                        c.number_format = "yyyy-mm-dd hh:mm"
    return buf.getvalue()
