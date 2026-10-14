"""M13 — P6 Relationship Converter.

Converts a P6 activity-layout export to Excel (sheet ``data``: one row per
activity, logic packed in the text columns ``Predecessor Details`` and
``Successor Details``) into a P6 spreadsheet-import workbook with the sheets
``TASK``, ``TASKPRED`` and ``Validation``.

Rules (all explicit, all reported):
- Rows with neither Activity Name nor WBS are export cell overflow and are dropped.
- Relationships are taken from Predecessor Details. Successor Details is used only
  to cross-check that both columns describe the same network.
- Lags and durations are written as exported (the layout's display unit, normally
  days). No unit conversion is performed.
- Activity Type is NOT in the layout; it is derived and declared as an assumption:
  OD = 0 and no Start -> Finish Milestone; OD = 0 and no Finish -> Start Milestone;
  otherwise Task Dependent.
- Start/Finish and Actual Duration are not exported to TASK: P6 recalculates them.
"""
from __future__ import annotations
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Mapping
import re

REQUIRED_COLUMNS = ("Activity ID", "Activity Status", "Activity Name", "WBS",
                    "Predecessor Details", "Successor Details")
TASK_COLUMNS = ("Orig Dur", "Rem Dur", "Start", "Finish", "Actual Start", "Actual Finish",
                "Primary Constraint", "Primary Constraint Date")
VALID_TYPES = ("FS", "SS", "FF", "SF")

TASK_FIELDS = ("task_code", "status_code", "task_name", "wbs_id", "task_type",
               "target_drtn_hr_cnt", "remain_drtn_hr_cnt", "act_start_date",
               "act_end_date", "cstr_type", "cstr_date")
TASK_LABELS = ("Activity ID", "Activity Status", "Activity Name", "WBS Code", "Activity Type",
               "Original Duration(d)", "Remaining Duration(d)", "Actual Start", "Actual Finish",
               "Primary Constraint", "Primary Constraint Date")
TASKPRED_FIELDS = ("Activity ID", "Activity Status", "Activity Name", "pred_task_id", "task_id",
                   "pred_type", "PREDTASK__status_code", "TASK__status_code", "lag_hr_cnt")
TASKPRED_LABELS = ("(*)Activity ID", "(*)Activity Status", "(*)Activity Name", "Predecessor",
                   "Successor", "Relationship Type", "(*)Predecessor Activity Status",
                   "(*)Successor Activity Status", "Lag(d)")

_TOKEN = re.compile(r"^(?P<id>.+?):\s*(?P<type>[A-Za-z]+)(?:\s+(?P<lag>-?\d+(?:\.\d+)?))?$")


class ConverterError(ValueError):
    pass


@dataclass(frozen=True)
class RelationshipToken:
    activity: str            # the other activity named in the token
    kind: str | None
    lag: float | None
    raw: str
    ok: bool


@dataclass(frozen=True)
class Issue:
    check: str
    activity_id: str | None
    detail: str


@dataclass(frozen=True)
class ConversionResult:
    task_rows: tuple[tuple[Any, ...], ...]
    taskpred_rows: tuple[tuple[Any, ...], ...]
    issues: tuple[Issue, ...]
    activity_count: int
    dropped_rows: int
    relationship_types: dict[str, int]
    lagged_relationships: int
    activity_types: dict[str, int]
    assumptions: tuple[str, ...] = field(default=())

    @property
    def relationship_count(self) -> int:
        return len(self.taskpred_rows)


def parse_relationship_details(text: Any) -> list[RelationshipToken]:
    """Split ``"A1980: FS, M1210: SS 57"`` into tokens. Unreadable tokens are kept with ok=False."""
    if text is None:
        return []
    tokens = []
    for raw in str(text).split(","):
        raw = raw.strip()
        if not raw:
            continue
        m = _TOKEN.match(raw)
        if not m:
            tokens.append(RelationshipToken(raw, None, None, raw, False))
            continue
        kind = m.group("type").upper()
        lag = float(m.group("lag")) if m.group("lag") else 0.0
        tokens.append(RelationshipToken(m.group("id").strip(), kind, lag, raw, kind in VALID_TYPES))
    return tokens


def read_data_sheet(source: bytes | str | Path, sheet: str = "data") -> list[dict[str, Any]]:
    """Read the export sheet into a list of dicts keyed by the header row."""
    from openpyxl import load_workbook
    stream = BytesIO(source) if isinstance(source, (bytes, bytearray)) else source
    try:
        wb = load_workbook(stream, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types
        raise ConverterError(f"The Excel file could not be opened: {exc}") from exc
    if sheet not in wb.sheetnames:
        raise ConverterError(f"The workbook has no sheet '{sheet}'. Sheets: {', '.join(wb.sheetnames)}")
    rows = wb[sheet].iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        raise ConverterError(f"Sheet '{sheet}' is empty")
    names = [str(h).strip() if h is not None else "" for h in header]
    out = [dict(zip(names, r)) for r in rows if r and any(v is not None for v in r)]
    wb.close()
    return out


def _num(value: Any) -> Any:
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _date(value: Any) -> datetime | None:
    return value if isinstance(value, datetime) else None


def _activity_type(row: Mapping[str, Any]) -> str:
    od = row.get("Orig Dur")
    if od in (0, 0.0):
        if row.get("Start") in (None, ""):
            return "Finish Milestone"
        if row.get("Finish") in (None, ""):
            return "Start Milestone"
    return "Task Dependent"


def convert_rows(rows: Iterable[Mapping[str, Any]]) -> ConversionResult:
    rows = list(rows)
    if not rows:
        raise ConverterError("The sheet has no data rows")
    missing = [c for c in REQUIRED_COLUMNS if c not in rows[0]]
    if missing:
        raise ConverterError("Missing required columns: " + ", ".join(missing))
    missing_task = [c for c in TASK_COLUMNS if c not in rows[0]]

    junk = [r for r in rows if r.get("Activity Name") in (None, "") and r.get("WBS") in (None, "")]
    data = [r for r in rows if not (r.get("Activity Name") in (None, "") and r.get("WBS") in (None, ""))]
    issues: list[Issue] = [Issue("Dropped row (export cell overflow)", r.get("Activity ID"),
                                 str(r.get("Predecessor Details") or "")) for r in junk]

    status: dict[str, Any] = {}
    for r in data:
        status.setdefault(r["Activity ID"], r.get("Activity Status"))
    for aid, n in Counter(r["Activity ID"] for r in data).items():
        if n > 1:
            issues.append(Issue("Duplicate Activity ID", aid, f"{n} rows"))

    # TASK
    task_rows = []
    types: Counter[str] = Counter()
    for r in data:
        t = _activity_type(r)
        types[t] += 1
        task_rows.append((r["Activity ID"], r.get("Activity Status"), r.get("Activity Name"), r.get("WBS"), t,
                          _num(r.get("Orig Dur")), _num(r.get("Rem Dur")), _date(r.get("Actual Start")),
                          _date(r.get("Actual Finish")), r.get("Primary Constraint") or None,
                          _date(r.get("Primary Constraint Date"))))

    # TASKPRED from Predecessor Details; cross-check with Successor Details
    pred_side: set[tuple] = set()
    succ_side: set[tuple] = set()
    seen: set[tuple] = set()
    taskpred = []
    rel_types: Counter[str] = Counter()
    lagged = 0
    for r in data:
        aid = r["Activity ID"]
        for tok in parse_relationship_details(r.get("Predecessor Details")):
            if not tok.ok:
                issues.append(Issue("Unreadable relationship (Predecessor Details)", aid, tok.raw))
                continue
            pred_side.add((tok.activity, aid, tok.kind, tok.lag))
            key = (tok.activity, aid, tok.kind)
            if key in seen:
                continue
            seen.add(key)
            rel_types[tok.kind] += 1
            lagged += tok.lag != 0
            taskpred.append((aid, r.get("Activity Status"), r.get("Activity Name"), tok.activity, aid,
                             tok.kind, status.get(tok.activity), r.get("Activity Status"), _num(tok.lag)))
        for tok in parse_relationship_details(r.get("Successor Details")):
            if not tok.ok:
                issues.append(Issue("Unreadable relationship (Successor Details)", aid, tok.raw))
                continue
            succ_side.add((aid, tok.activity, tok.kind, tok.lag))

    fmt = lambda k: f"{k[0]} -> {k[1]} {k[2]} {_num(k[3])}"
    for k in sorted(pred_side - succ_side):
        issues.append(Issue("Relationship only in Predecessor Details", k[1], fmt(k)))
    for k in sorted(succ_side - pred_side):
        issues.append(Issue("Relationship only in Successor Details", k[0], fmt(k)))
    for k in sorted(pred_side):
        if k[0] not in status:
            issues.append(Issue("Predecessor not found in data", k[1], fmt(k)))
    for k in sorted(succ_side):
        if k[1] not in status:
            issues.append(Issue("Successor not found in data", k[0], fmt(k)))
    for k in sorted(pred_side):
        if k[0] == k[1]:
            issues.append(Issue("Self-relationship", k[0], fmt(k)))

    assumptions = [
        "Activity Type derived: OD = 0 and no Start -> Finish Milestone; OD = 0 and no Finish -> "
        "Start Milestone; otherwise Task Dependent. LOE / WBS Summary cannot be detected.",
        "Durations and lags written as exported (layout display unit, normally days); the P6 user "
        "importing must use the same time unit.",
        "Start, Finish and Actual Duration are not imported: P6 recalculates them.",
        "wbs_id is written exactly as exported; WBS elements must exist in the target project.",
    ]
    if missing_task:
        assumptions.append("Columns missing for TASK (left blank): " + ", ".join(missing_task))
    return ConversionResult(tuple(task_rows), tuple(taskpred), tuple(issues), len(data), len(junk),
                            dict(rel_types), lagged, dict(types), tuple(assumptions))


def convert_workbook(source: bytes | str | Path, sheet: str = "data") -> ConversionResult:
    return convert_rows(read_data_sheet(source, sheet))


def to_import_workbook(result: ConversionResult) -> bytes:
    """Build the P6 import workbook: row 1 field names, row 2 titles, data from row 3."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    yellow = PatternFill("solid", fgColor="FFFF00")
    bold = Font(bold=True)
    wb = Workbook()

    def sheet(ws, fields, labels, rows, widths, date_cols=()):
        ws.append(list(fields))
        ws.append(list(labels))
        for c in ws[1]:
            c.font = bold
        for c in ws[2]:
            c.fill = yellow
        for r in rows:
            ws.append(list(r))
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        for col in date_cols:
            for cell in ws[col][2:]:
                cell.number_format = "dd-mmm-yy"
        ws.freeze_panes = "A3"

    ws = wb.active
    ws.title = "TASK"
    sheet(ws, TASK_FIELDS, TASK_LABELS, result.task_rows,
          (18, 14, 60, 30, 17, 10, 10, 12, 12, 20, 12), ("H", "I", "K"))
    sheet(wb.create_sheet("TASKPRED"), TASKPRED_FIELDS, TASKPRED_LABELS, result.taskpred_rows,
          (18, 14, 50, 18, 18, 8, 16, 16, 8))
    v = wb.create_sheet("Validation")
    v.append(["Check", "Activity ID", "Detail"])
    for c in v[1]:
        c.font = bold
    v.append(["Summary", None, f"Activities: {result.activity_count}"])
    v.append(["Summary", None, f"Relationships (TASKPRED): {result.relationship_count}"])
    v.append(["Summary", None, "No issues found" if not result.issues
              else f"{len(result.issues)} issues (see below)"])
    for i in result.issues:
        v.append([i.check, i.activity_id, i.detail])
    for a in result.assumptions:
        v.append(["Assumption", None, a])
    v.column_dimensions["A"].width = 48
    v.column_dimensions["B"].width = 18
    v.column_dimensions["C"].width = 90
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
